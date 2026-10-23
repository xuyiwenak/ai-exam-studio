"""Convert local source workbooks into private exam data served by the local app."""

import json
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(__file__).with_name("source")
OUTPUT = Path(__file__).with_name(".local") / "questions.js"


def clean(value):
    return "" if value is None else str(value).strip()


def vibe_questions():
    sheet = load_workbook(SOURCE / "VibeCoding.xlsx", read_only=True, data_only=True).active
    questions = []
    for row_number, row in enumerate(sheet.values, start=1):
        if row_number == 1 or not clean(row[2]):
            continue
        kind = {"单选题": "single", "多选题": "multiple", "是非题": "boolean", "简答题": "essay"}[clean(row[0])]
        options = [clean(value) for value in row[3:10] if clean(value)]
        answer = [] if kind == "essay" else [int(char) - 1 for char in clean(row[10])]
        assert all(0 <= index < len(options) for index in answer), (row_number, answer)
        questions.append({
            "id": f"vibe-{row_number}",
            "type": kind,
            "text": clean(row[2]),
            "options": options,
            "answer": answer,
            "explanation": clean(row[11]),
        })
    return questions


def opc_questions():
    sheet = load_workbook(SOURCE / "OPC题库.xlsx", read_only=True, data_only=True).active
    questions = []
    for row_number, row in enumerate(sheet.values, start=1):
        if row_number == 1 or not clean(row[2]):
            continue
        kind = {"单选": "single", "多选": "multiple", "判断": "boolean"}[clean(row[1])]
        options = [clean(value) for value in row[4:9] if clean(value)]
        if kind == "boolean":
            options = ["正确", "错误"]
            answer = [options.index(clean(row[3]))]
        else:
            answer = [ord(part) - ord("A") for part in clean(row[3]).split("|")]
        invalid = any(index < 0 or index >= len(options) for index in answer)
        questions.append({
            "id": f"opc-{row_number}",
            "type": kind,
            "text": clean(row[2]),
            "options": options,
            "answer": answer,
            "explanation": "",
            "invalid": invalid,
        })
    return questions


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    banks = {
        "vibe": {"title": "VibeCoding", "source": "VibeCoding.xlsx", "questions": vibe_questions()},
        "opc": {"title": "OPC 智能体应用师", "source": "OPC题库.xlsx", "questions": opc_questions()},
    }
    assert len(banks["vibe"]["questions"]) == 214
    assert len(banks["opc"]["questions"]) == 200
    OUTPUT.write_text("window.EXAM_BANKS = " + json.dumps(banks, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(f"Wrote {OUTPUT} with 414 questions")


if __name__ == "__main__":
    main()
